#!/usr/bin/env python3
"""
verify_rich.py — Phase 3 Rich Mode (Sprint 3, light stub)

Provides a single real check (`verify_xlsx()`) that loads the generated
`22_Traceability_Matrix.xlsx` and prints sheet names + row counts. Sprint 5
will add the full deep-enrichment validation (17-field schema, FR/NFR/Risk
trace consistency).

Usage:
  python3 scripts/verify_rich.py
  python3 scripts/verify_rich.py --xlsx path/to/22_Traceability_Matrix.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List


def verify_xlsx(xlsx_path: Path) -> List[dict]:
    """Load xlsx and print/return sheet info. Returns list of dicts."""
    try:
        import openpyxl
    except ImportError:
        print("[err] openpyxl not installed; pip install openpyxl", file=sys.stderr)
        return []

    if not xlsx_path.is_file():
        print(f"[err] xlsx not found: {xlsx_path}", file=sys.stderr)
        return []

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    rows: List[dict] = []
    print(f"[verify_xlsx] {xlsx_path}")
    print(f"[verify_xlsx] sheets={len(wb.sheetnames)}")
    for s in wb.sheetnames:
        ws = wb[s]
        info = {"sheet": s, "rows": ws.max_row, "cols": ws.max_column}
        rows.append(info)
        print(f"  - {s}: rows={info['rows']} cols={info['cols']}")
    total = sum(r["rows"] for r in rows)
    print(f"[verify_xlsx] total_rows={total}")
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description="Verify Phase 3 RICH artefacts (Sprint 3 light)")
    parser.add_argument("--xlsx", default=None,
                        help="Path to 22_Traceability_Matrix.xlsx (default: ../22_Traceability_Matrix.xlsx)")
    args = parser.parse_args()

    base = Path(__file__).resolve().parent.parent
    xlsx = Path(args.xlsx) if args.xlsx else base / "22_Traceability_Matrix.xlsx"
    rows = verify_xlsx(xlsx)
    if not rows:
        return 1
    print("[ok] verify_xlsx PASS")
    print("[note] full DEEP validation arrives in Sprint 5 (17-field schema, FR/NFR/Risk traces)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
