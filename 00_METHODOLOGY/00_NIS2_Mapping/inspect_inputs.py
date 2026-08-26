"""Inspect all 3 NIST xlsx files in detail."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
from pathlib import Path

base = Path(r"C:\Users\paulo.ribeiro\OneDrive - Instituto CCG ZGDV\Ambiente de Trabalho\Methodology-main\Methodology-main\00_METHODOLOGY\PREPROCESSING_by_domain")

for name in [
    "NIST-Privacy-Framework-V1.0-Core.xlsx",
    "nist_ai_rmf_playbook.xlsx",
    "PF 1.0 and 1.1_Core Mapping.xlsx",
]:
    path = base / name
    print(f"\n{'='*80}\n{name}  ({path.stat().st_size:,} bytes)\n{'='*80}")
    wb = load_workbook(path, data_only=True, read_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n  Sheet: {sname!r}")
        # Show first 6 rows of content (max 6 cols)
        rows_shown = 0
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            cells = [str(c)[:50] if c is not None else "" for c in row[:6]]
            print(f"    R{i}: {cells}")
            rows_shown += 1
            if rows_shown >= 6:
                break
